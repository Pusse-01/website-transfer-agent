"""
Excel reader module for loading page lists from spreadsheets.

Supports two Excel formats:
1. Blog Post List - columns: Priority, ID, Categories, Title, Status, Url Key, Published At
2. Static Page List - multiple sheets (1-General, 4-Categories, 5-Delivery, etc.)
   with columns: A (row#), Title, Priority, Done, Review, Remark,
   Builder URL (Traditional Chinese), Original Page (Traditional Chinese),
   Simplified Chinese, English
"""

import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def read_blog_list(excel_path: str, sheet_name: str = None) -> list[dict]:
    """
    Read blog post list from an Excel file.
    Expected columns: Priority, ID, Categories, Title, Status, Url Key

    Only returns posts with Status == 'Published'.

    Returns:
        List of dicts with keys: priority, id, categories, title, status, url_key, published_at
    """
    try:
        import openpyxl
    except ImportError:
        logger.error("openpyxl is required: pip install openpyxl")
        return []

    path = Path(excel_path)
    if not path.exists():
        logger.error(f"Excel file not found: {excel_path}")
        return []

    try:
        wb = openpyxl.load_workbook(path, read_only=True)

        sheet = None
        if sheet_name and sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
        else:
            # Try to find the blog post list sheet
            for name in wb.sheetnames:
                if "blog" in name.lower() or "post" in name.lower():
                    sheet = wb[name]
                    break

            if sheet is None:
                if len(wb.sheetnames) > 1:
                    sheet = wb[wb.sheetnames[0]]
                else:
                    sheet = wb.active

        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            wb.close()
            return []

        # Find header row
        header_row = rows[0]
        headers = [str(h).strip().lower().replace(" ", "_") if h else "" for h in header_row]

        # Map column indices
        col_map = {}
        for i, h in enumerate(headers):
            if "priority" in h:
                col_map["priority"] = i
            elif h == "id":
                col_map["id"] = i
            elif "categor" in h:
                col_map["categories"] = i
            elif "title" in h:
                col_map["title"] = i
            elif "status" in h:
                col_map["status"] = i
            elif "url" in h and "key" in h:
                col_map["url_key"] = i
            elif "url_key" in h:
                col_map["url_key"] = i
            elif "published" in h:
                col_map["published_at"] = i

        if "url_key" not in col_map:
            logger.error(f"Could not find 'Url Key' column in sheet. Headers: {headers}")
            wb.close()
            return []

        blog_posts = []
        for row in rows[1:]:
            if not row or all(cell is None for cell in row):
                continue

            url_key = row[col_map["url_key"]] if col_map.get("url_key") is not None else None
            if not url_key:
                continue

            status = str(row[col_map["status"]]).strip() if col_map.get("status") is not None and row[col_map["status"]] else ""

            # Only include published posts
            if status.lower() != "published":
                continue

            post = {
                "priority": row[col_map["priority"]] if col_map.get("priority") is not None else 0,
                "id": row[col_map["id"]] if col_map.get("id") is not None else "",
                "categories": row[col_map["categories"]] if col_map.get("categories") is not None else "",
                "title": row[col_map["title"]] if col_map.get("title") is not None else "",
                "status": status,
                "url_key": str(url_key).strip(),
                "published_at": str(row[col_map["published_at"]]) if col_map.get("published_at") is not None and row[col_map["published_at"]] else "",
                "page_type": "blog",
            }
            blog_posts.append(post)

        wb.close()
        logger.info(f"Loaded {len(blog_posts)} published blog posts from {excel_path}")
        return blog_posts

    except Exception as e:
        logger.error(f"Failed to read Excel file: {e}")
        return []


def read_static_page_list(excel_path: str) -> list[dict]:
    """
    Read static page list from an Excel file with multiple category sheets.

    The static page Excel has sheets like: Introduction, 1-General, 4-Categories,
    5-Delivery, Brand Collection, Campaign, 6New Housing, Tutorial.

    Each sheet (except Introduction) has pages with columns:
    - A: row number / page ID
    - B: Title
    - C: Priority (some sheets)
    - D: Done
    - E: Review
    - F: Remark
    - G: Builder URL (Traditional Chinese)
    - H: Original Page (Traditional Chinese) - contains title + URL on alternate rows
    - I: Simplified Chinese - title + URL
    - J: English - title + URL

    Returns:
        List of dicts with keys: id, title, category, page_type,
        urls (dict with zh_hk, zh_cn, en variants), done, review, remark, priority
    """
    try:
        import openpyxl
    except ImportError:
        logger.error("openpyxl is required: pip install openpyxl")
        return []

    path = Path(excel_path)
    if not path.exists():
        logger.error(f"Excel file not found: {excel_path}")
        return []

    try:
        wb = openpyxl.load_workbook(path, read_only=True)

        # Skip 'Introduction' sheet, process all others
        skip_sheets = {"introduction"}
        static_pages = []

        for sheet_name in wb.sheetnames:
            if sheet_name.lower() in skip_sheets:
                continue

            sheet = wb[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue

            # Detect header row - look for row containing 'Title'
            header_idx = 0
            headers = []
            for idx, row in enumerate(rows):
                row_strs = [str(c).strip().lower() if c else "" for c in row]
                if any("title" in s for s in row_strs):
                    header_idx = idx
                    headers = row_strs
                    break

            if not headers:
                # Try first row as header
                headers = [str(c).strip().lower() if c else "" for c in rows[0]]
                header_idx = 0

            # Map columns by header content
            col_map = {}
            for i, h in enumerate(headers):
                if "title" in h and "title" not in col_map:
                    col_map["title"] = i
                elif "priority" in h:
                    col_map["priority"] = i
                elif "done" in h:
                    col_map["done"] = i
                elif "review" in h:
                    col_map["review"] = i
                elif "remark" in h:
                    col_map["remark"] = i
                elif "builder" in h:
                    col_map["builder_url"] = i
                elif "original" in h and ("traditional" in h or "chinese" in h or "page" in h):
                    col_map["zh_hk"] = i
                elif "simplified" in h:
                    col_map["zh_cn"] = i
                elif "english" in h:
                    col_map["en"] = i

            # The first column (A) typically has the page number/ID
            id_col = 0

            # Process data rows - static pages often have title on one row
            # and URL on the next row for each language variant
            data_rows = rows[header_idx + 1:]

            # Group rows by page entry - pages are identified by having a value in column A
            i = 0
            while i < len(data_rows):
                row = data_rows[i]
                if not row or all(cell is None for cell in row):
                    i += 1
                    continue

                page_id = row[id_col] if row[id_col] is not None else None
                if page_id is None:
                    i += 1
                    continue

                title = str(row[col_map["title"]]).strip() if col_map.get("title") is not None and row[col_map["title"]] else ""
                if not title:
                    i += 1
                    continue

                # Collect URLs from this row and the next row
                urls = {"zh_hk": [], "zh_cn": [], "en": []}

                # Check current row and next row for URL data
                rows_to_check = [row]
                if i + 1 < len(data_rows):
                    next_row = data_rows[i + 1]
                    # Next row belongs to same entry if column A is empty
                    if next_row and (next_row[id_col] is None or str(next_row[id_col]).strip() == ""):
                        rows_to_check.append(next_row)
                        i += 1  # Skip next row since we're consuming it

                for check_row in rows_to_check:
                    for lang, col_key in [("zh_hk", "zh_hk"), ("zh_cn", "zh_cn"), ("en", "en")]:
                        if col_key in col_map and check_row[col_map[col_key]]:
                            val = str(check_row[col_map[col_key]]).strip()
                            if val and val.lower() != "false" and val.lower() != "none":
                                urls[lang].append(val)

                # Extract URL (typically the one starting with http)
                url_data = {}
                for lang in ["zh_hk", "zh_cn", "en"]:
                    page_url = ""
                    page_title = ""
                    for val in urls[lang]:
                        if val.startswith("http"):
                            page_url = val
                        else:
                            page_title = val
                    url_data[lang] = {"url": page_url, "title": page_title}

                # Extract the primary URL (Traditional Chinese is the main one)
                primary_url = url_data.get("zh_hk", {}).get("url", "")
                # Also check the builder URL column
                builder_url = ""
                if col_map.get("builder_url") is not None:
                    for check_row in rows_to_check:
                        val = check_row[col_map["builder_url"]]
                        if val and str(val).strip().startswith("http"):
                            builder_url = str(val).strip()
                            break

                # Extract URL key from the primary URL
                url_key = _extract_url_key_from_url(primary_url)

                page = {
                    "id": page_id,
                    "title": title,
                    "category": sheet_name,
                    "page_type": "static",
                    "url_key": url_key,
                    "primary_url": primary_url,
                    "builder_url": builder_url,
                    "urls": url_data,
                    "done": str(row[col_map["done"]]).strip() if col_map.get("done") is not None and row[col_map["done"]] else "",
                    "review": str(row[col_map["review"]]).strip() if col_map.get("review") is not None and row[col_map["review"]] else "",
                    "remark": str(row[col_map["remark"]]).strip() if col_map.get("remark") is not None and row[col_map["remark"]] else "",
                    "priority": str(row[col_map["priority"]]).strip() if col_map.get("priority") is not None and row[col_map["priority"]] else "",
                }
                static_pages.append(page)

                i += 1

        wb.close()
        logger.info(f"Loaded {len(static_pages)} static pages from {excel_path}")
        return static_pages

    except Exception as e:
        logger.error(f"Failed to read static page Excel file: {e}", exc_info=True)
        return []


def read_tutorial_list(excel_path: str) -> list[dict]:
    """
    Read Tutorial sheet from the static page Excel.
    Tutorial sheet has a simpler structure:
    A: number, B: Title, C: Original Page URL, D: Simplified Chinese, E: English

    Returns:
        List of dicts with same structure as read_static_page_list
    """
    try:
        import openpyxl
    except ImportError:
        logger.error("openpyxl is required: pip install openpyxl")
        return []

    path = Path(excel_path)
    if not path.exists():
        return []

    try:
        wb = openpyxl.load_workbook(path, read_only=True)

        if "Tutorial" not in wb.sheetnames:
            wb.close()
            return []

        sheet = wb["Tutorial"]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            wb.close()
            return []

        # Find header
        header_idx = 0
        for idx, row in enumerate(rows):
            row_strs = [str(c).strip().lower() if c else "" for c in row]
            if any("title" in s for s in row_strs):
                header_idx = idx
                break

        headers = [str(c).strip().lower() if c else "" for c in rows[header_idx]]

        col_map = {}
        for i, h in enumerate(headers):
            if "title" in h and "title" not in col_map:
                col_map["title"] = i
            elif "original" in h:
                col_map["zh_hk"] = i
            elif "simplified" in h:
                col_map["zh_cn"] = i
            elif "english" in h:
                col_map["en"] = i

        tutorials = []
        data_rows = rows[header_idx + 1:]

        i = 0
        while i < len(data_rows):
            row = data_rows[i]
            if not row or all(cell is None for cell in row):
                i += 1
                continue

            page_id = row[0]
            if page_id is None:
                i += 1
                continue

            title = str(row[col_map["title"]]).strip() if col_map.get("title") is not None and row[col_map["title"]] else ""
            if not title:
                i += 1
                continue

            # Collect URLs from this row and possibly next row
            urls = {"zh_hk": [], "zh_cn": [], "en": []}
            rows_to_check = [row]
            if i + 1 < len(data_rows):
                next_row = data_rows[i + 1]
                if next_row and (next_row[0] is None or str(next_row[0]).strip() == ""):
                    rows_to_check.append(next_row)
                    i += 1

            for check_row in rows_to_check:
                for lang, col_key in [("zh_hk", "zh_hk"), ("zh_cn", "zh_cn"), ("en", "en")]:
                    if col_key in col_map and len(check_row) > col_map[col_key] and check_row[col_map[col_key]]:
                        val = str(check_row[col_map[col_key]]).strip()
                        if val and val.lower() != "false" and val.lower() != "none":
                            urls[lang].append(val)

            url_data = {}
            for lang in ["zh_hk", "zh_cn", "en"]:
                page_url = ""
                page_title = ""
                for val in urls[lang]:
                    if val.startswith("http"):
                        page_url = val
                    else:
                        page_title = val
                url_data[lang] = {"url": page_url, "title": page_title}

            primary_url = url_data.get("zh_hk", {}).get("url", "")
            url_key = _extract_url_key_from_url(primary_url)

            tutorials.append({
                "id": page_id,
                "title": title,
                "category": "Tutorial",
                "page_type": "static",
                "url_key": url_key,
                "primary_url": primary_url,
                "builder_url": "",
                "urls": url_data,
                "done": "",
                "review": "",
                "remark": "",
                "priority": "",
            })

            i += 1

        wb.close()
        logger.info(f"Loaded {len(tutorials)} tutorial pages")
        return tutorials

    except Exception as e:
        logger.error(f"Failed to read Tutorial sheet: {e}", exc_info=True)
        return []


def detect_excel_type(excel_path: str) -> str:
    """
    Detect whether an Excel file is a Blog Post List or Static Page List.

    Returns:
        'blog' if it's a blog post list
        'static' if it's a static page list
        'unknown' if can't determine
    """
    try:
        import openpyxl
    except ImportError:
        return "unknown"

    path = Path(excel_path)
    if not path.exists():
        return "unknown"

    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        sheet_names = [s.lower() for s in wb.sheetnames]

        # Blog post list typically has "Blog post list" sheet or columns like Url Key, Status
        is_blog = any("blog" in s or "post" in s for s in sheet_names)

        # Static page list has sheets like "1-General", "4-Categories", "Tutorial"
        is_static = any(
            s.startswith("1-") or "general" in s or "categor" in s
            or "delivery" in s or "tutorial" in s or "campaign" in s
            for s in sheet_names
        )

        wb.close()

        if is_blog and not is_static:
            return "blog"
        elif is_static and not is_blog:
            return "static"
        elif is_blog:
            return "blog"
        elif is_static:
            return "static"
        else:
            # Check first sheet headers
            wb = openpyxl.load_workbook(path, read_only=True)
            sheet = wb.active
            rows = list(sheet.iter_rows(max_row=1, values_only=True))
            wb.close()
            if rows:
                headers = [str(h).lower() for h in rows[0] if h]
                if any("url_key" in h or "url key" in h for h in headers):
                    return "blog"
                if any("builder" in h for h in headers):
                    return "static"
            return "unknown"

    except Exception:
        return "unknown"


def _extract_url_key_from_url(url: str) -> str:
    """Extract a URL key/identifier from a full URL.

    Examples:
        https://www.pricerite.com.hk/hk/zh/member-point -> member-point
        https://www.pricerite.com.hk/hk/zh/clear-cache-and-cookies -> clear-cache-and-cookies
        https://www.pricerite.com.hk/mattress-eform -> mattress-eform
    """
    if not url:
        return ""

    from urllib.parse import urlparse
    parsed = urlparse(url.strip())
    path = parsed.path.rstrip("/")

    if not path:
        return ""

    # Get the last segment of the path
    segments = [s for s in path.split("/") if s]
    if not segments:
        return ""

    # The URL key is the last meaningful segment
    url_key = segments[-1]

    # Skip language prefixes if the last segment is a lang code
    if url_key in ("zh", "en", "sc", "hk", "mo") and len(segments) > 1:
        url_key = segments[-2] if len(segments) >= 2 else url_key

    return url_key
